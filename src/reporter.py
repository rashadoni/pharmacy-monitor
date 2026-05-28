"""Генерация отчёта: HTML-письмо + Excel-вложение."""

from __future__ import annotations

from dataclasses import asdict
from io import BytesIO
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.analyzer import AnalysisReport

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
_env = Environment(
    loader=FileSystemLoader(TEMPLATES_DIR),
    autoescape=select_autoescape(["html", "xml"]),
)


def render_html(report: AnalysisReport) -> str:
    template = _env.get_template("report.html.j2")
    return template.render(**asdict(report))


def render_excel(report: AnalysisReport) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1d1d1f")
    header_align = Alignment(horizontal="left", vertical="center")

    def add_sheet(name: str, headers: list[str], rows: list[list]) -> None:
        ws = wb.create_sheet(title=name[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    # Undercuts
    add_sheet(
        "Конкурент дешевле",
        [
            "Товар",
            "Клиент ₼",
            "Клиент URL",
            "Конкурент сайт",
            "Конкурент ₼",
            "Конкурент URL",
            "Дельта %",
        ],
        [
            [
                u.canonical_name,
                u.client_price,
                u.client_url,
                u.competitor_site,
                u.competitor_price,
                u.competitor_url,
                u.diff_pct,
            ]
            for u in report.undercuts
        ],
    )

    add_sheet(
        "Изменения цен",
        ["Сайт", "Товар", "URL", "Было ₼", "Стало ₼", "Дельта %", "Промо"],
        [
            [
                c.site,
                c.product_name,
                c.url,
                c.prev_price,
                c.curr_price,
                c.delta_pct,
                c.promo_label or "",
            ]
            for c in report.price_changes
        ],
    )

    add_sheet(
        "Новые SKU",
        ["Сайт", "Категория", "Товар", "URL", "Цена ₼"],
        [[n.site, n.category or "", n.name, n.url, n.price] for n in report.new_products],
    )

    add_sheet(
        "Промо",
        ["Сайт", "Заголовок", "Landing URL", "Статус"],
        [
            [p.site, p.title, p.landing_url or "", "новая" if p.is_new else "завершилась"]
            for p in report.promo_changes
        ],
    )

    if not wb.sheetnames:
        ws = wb.create_sheet("Empty")
        ws["A1"] = "Нет изменений с предыдущего прогона."

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_filename(report: AnalysisReport) -> str:
    return f"pharmacy-monitor-{report.run_started_at.strftime('%Y-%m-%d')}.xlsx"


def email_subject(report: AnalysisReport) -> str:
    parts = []
    if report.summary_counts.get("undercuts", 0) > 0:
        parts.append(f"⚠️ {report.summary_counts['undercuts']} undercuts")
    if report.summary_counts.get("price_changes", 0) > 0:
        parts.append(f"{report.summary_counts['price_changes']} price changes")
    if report.summary_counts.get("new_products", 0) > 0:
        parts.append(f"{report.summary_counts['new_products']} new SKU")
    suffix = " — " + ", ".join(parts) if parts else ""
    return f"Pharmacy Monitor {report.run_started_at.strftime('%d.%m.%Y')}{suffix}"
