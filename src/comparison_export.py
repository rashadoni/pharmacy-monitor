"""Excel-выгрузка страницы сравнения: «все товары с разной ценой» одним файлом.

Запрос клиента (2026-10): «fərqli qiymətə olan bütün məhsulları Excel faylına
download etmək imkanı». До этого на странице была только CSV-выгрузка, собранная
в браузере: без BOM (азербайджанские буквы в Excel превращались в кракозябры),
с запятой-разделителем (в az/ru-локали Excel кладёт всю строку в одну ячейку) и
с ценами-текстом. Настоящий .xlsx открывается сразу таблицей, с числами,
фильтрами и рабочими ссылками.

Лист 1 — только строки, где цены на сайтах РАЗНЫЕ (то, что просили). Лист 2 —
вся текущая выборка страницы, чтобы файл можно было переслать без оговорок
«а где остальное». Лист 3 — когда и с какими фильтрами выгружено.

Строки приходят из `api._comparison_rows` — те же, что видит страница; здесь
нет своей логики цен, только оформление.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SITES = ("pharmonline", "aptekonline", "aloe")
CLIENT_SITE = "pharmonline"
SITE_TITLES = {"pharmonline": "Pharmonline", "aptekonline": "Aptekonline", "aloe": "Aloe"}
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Время в файле — бакинское: файл читает клиент, а не сервер. В Азербайджане
# нет перехода на летнее время, поэтому фиксированного сдвига достаточно.
_BAKU_OFFSET = timedelta(hours=4)

_STRINGS: dict[str, dict[str, str]] = {
    "az": {
        "sheet_diff": "Fərqli qiymətlər",
        "sheet_all": "Bütün müqayisə",
        "sheet_info": "Məlumat",
        "name": "Məhsul",
        "brand": "Brend",
        "price": "{site}, AZN",
        "cheapest": "Ən ucuz",
        "diff_azn": "Fərq, AZN",
        "diff_pct": "Fərq, %",
        "position": "Pharmonline mövqeyi",
        "note": "Qeyd",
        "link": "{site} keçid",
        "pos_cheapest": "Ən ucuz",
        "pos_priciest": "Ən baha",
        "pos_middle": "Ortada",
        "pos_equal": "Eyni qiymət",
        "note_unit": "Qablaşdırma fərqlidir — 1 ədədin qiyməti müqayisə olunur: {parts}",
        "note_unit_part": "{site} {price} ({count} əd.)",
        "note_stale": "{site}: qiymət {days} gün əvvəlkidir, fərqə daxil edilməyib",
        "note_review": "Fərq çox böyükdür — məhsulların eyni olduğunu yoxlayın",
        "info_title": "Qiymət müqayisəsi",
        "info_generated": "Fayl hazırlanıb",
        "info_search": "Axtarış",
        "info_category": "Kateqoriya",
        "info_min_sites": "Ən azı neçə saytda",
        "info_aloe": "Yalnız Aloe-də olanlar",
        "info_diff_rows": "Fərqli qiymətli məhsullar",
        "info_all_rows": "Seçimdəki bütün məhsullar",
        "info_rule": "«Fərqli qiymətlər» vərəqi: saytlarda qiyməti eyni olmayan bütün məhsullar, "
        "ən böyük fərqdən başlayaraq.",
        "info_colors": "Yaşıl — ən aşağı qiymət, qırmızı — ən yüksək.",
        "yes": "bəli",
        "all": "hamısı",
        "filename": "qiymet-ferqleri",
    },
    "ru": {
        "sheet_diff": "Разные цены",
        "sheet_all": "Всё сравнение",
        "sheet_info": "О файле",
        "name": "Товар",
        "brand": "Бренд",
        "price": "{site}, AZN",
        "cheapest": "Дешевле всего",
        "diff_azn": "Разница, AZN",
        "diff_pct": "Разница, %",
        "position": "Позиция Pharmonline",
        "note": "Примечание",
        "link": "Ссылка {site}",
        "pos_cheapest": "Самая низкая",
        "pos_priciest": "Самая высокая",
        "pos_middle": "Посередине",
        "pos_equal": "Цена одинаковая",
        "note_unit": "Фасовка разная — сравнивается цена за 1 шт.: {parts}",
        "note_unit_part": "{site} {price} ({count} шт.)",
        "note_stale": "{site}: цене {days} дн., в разницу не входит",
        "note_review": "Разница очень большая — проверьте, что товары одинаковые",
        "info_title": "Сравнение цен",
        "info_generated": "Файл сформирован",
        "info_search": "Поиск",
        "info_category": "Категория",
        "info_min_sites": "Минимум сайтов",
        "info_aloe": "Только те, что есть на Aloe",
        "info_diff_rows": "Товаров с разной ценой",
        "info_all_rows": "Всего товаров в выборке",
        "info_rule": "Лист «Разные цены»: все товары, у которых цена на сайтах не совпадает, "
        "начиная с самой большой разницы.",
        "info_colors": "Зелёным — самая низкая цена, красным — самая высокая.",
        "yes": "да",
        "all": "все",
        "filename": "raznye-ceny",
    },
    "en": {
        "sheet_diff": "Price differences",
        "sheet_all": "Full comparison",
        "sheet_info": "About",
        "name": "Product",
        "brand": "Brand",
        "price": "{site}, AZN",
        "cheapest": "Cheapest",
        "diff_azn": "Difference, AZN",
        "diff_pct": "Difference, %",
        "position": "Pharmonline position",
        "note": "Note",
        "link": "{site} link",
        "pos_cheapest": "Lowest",
        "pos_priciest": "Highest",
        "pos_middle": "In between",
        "pos_equal": "Same price",
        "note_unit": "Pack sizes differ — price per unit is compared: {parts}",
        "note_unit_part": "{site} {price} ({count} pcs)",
        "note_stale": "{site}: price is {days} days old, not counted in the difference",
        "note_review": "The difference is very large — check that the products are the same",
        "info_title": "Price comparison",
        "info_generated": "Generated",
        "info_search": "Search",
        "info_category": "Category",
        "info_min_sites": "Minimum sites",
        "info_aloe": "Only products listed on Aloe",
        "info_diff_rows": "Products with different prices",
        "info_all_rows": "Products in this selection",
        "info_rule": "Sheet “Price differences”: every product whose price is not the same "
        "across sites, largest difference first.",
        "info_colors": "Green — lowest price, red — highest.",
        "yes": "yes",
        "all": "all",
        "filename": "price-differences",
    },
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
_LOW_FONT = Font(color="15803D", bold=True)
_HIGH_FONT = Font(color="B91C1C")
_STALE_FONT = Font(color="9CA3AF", strike=True)
_LINK_FONT = Font(color="1D4ED8", underline="single")
_MONEY = "0.00"
_PERCENT = "0.0%"


def strings(locale: str) -> dict[str, str]:
    return _STRINGS.get(locale, _STRINGS["az"])


def has_price_difference(row: dict[str, Any]) -> bool:
    """Цены на сайтах сравнимы и НЕ равны.

    `spread_pct` округлён до десятых, поэтому «10.00 против 10.03» дал бы 0.3% →
    при пороге можно потерять строку; сравниваем сами min/max.
    """
    return (
        row.get("spread_pct") is not None
        and row.get("min_price") is not None
        and row["min_price"] != row["max_price"]
    )


def filename(locale: str, generated_at: datetime) -> str:
    local = generated_at + _BAKU_OFFSET
    return f"{strings(locale)['filename']}-{local:%Y-%m-%d}.xlsx"


def _write_text(cell: Any, value: Any) -> None:
    """Записать текст с чужого сайта так, чтобы он остался текстом.

    Управляющие символы openpyxl отвергает исключением. Строку, начинающуюся с
    «=», он записал бы как формулу — возвращаем ей строковый тип. (Ведущие «+»,
    «-», «@» опасны только в CSV: в .xlsx у ячейки есть тип, и строка не
    вычисляется.)
    """
    if value is None or value == "":
        return
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    if cell.data_type == "f":
        cell.data_type = "s"


def _compared_value(row: dict[str, Any], entry: dict[str, Any]) -> float:
    if row.get("spread_basis") == "unit" and entry.get("unit_price") is not None:
        return entry["unit_price"]
    return entry["price"]


def _position(row: dict[str, Any], t: dict[str, str]) -> str:
    entry = row["prices"].get(CLIENT_SITE)
    if entry is None or entry.get("stale") or row.get("min_price") is None:
        return ""
    if row["min_price"] == row["max_price"]:
        return t["pos_equal"]
    value = _compared_value(row, entry)
    if value == row["min_price"]:
        return t["pos_cheapest"]
    if value == row["max_price"]:
        return t["pos_priciest"]
    return t["pos_middle"]


def _cheapest_sites(row: dict[str, Any]) -> str:
    if row.get("min_price") is None:
        return ""
    return ", ".join(
        SITE_TITLES[site]
        for site in SITES
        if (entry := row["prices"].get(site)) is not None
        and not entry.get("stale")
        and _compared_value(row, entry) == row["min_price"]
    )


def _note(row: dict[str, Any], t: dict[str, str]) -> str:
    notes: list[str] = []
    if row.get("spread_basis") == "unit":
        parts = [
            t["note_unit_part"].format(
                site=SITE_TITLES[site],
                price=f"{entry['unit_price']:.2f}",
                count=entry.get("pack_count") or 1,
            )
            for site in SITES
            if (entry := row["prices"].get(site)) is not None and not entry.get("stale")
        ]
        notes.append(t["note_unit"].format(parts="; ".join(parts)))
    for site in SITES:
        entry = row["prices"].get(site)
        if entry is not None and entry.get("stale") and entry.get("age_days") is not None:
            notes.append(t["note_stale"].format(site=SITE_TITLES[site], days=entry["age_days"]))
    if row.get("needs_review"):
        notes.append(t["note_review"])
    return ". ".join(notes)


def _write_sheet(ws: Any, rows: list[dict[str, Any]], t: dict[str, str]) -> None:
    headers = [
        t["name"],
        t["brand"],
        *(t["price"].format(site=SITE_TITLES[s]) for s in SITES),
        t["cheapest"],
        t["diff_azn"],
        t["diff_pct"],
        t["position"],
        t["note"],
        *(t["link"].format(site=SITE_TITLES[s]) for s in SITES),
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    ws.row_dimensions[1].height = 32

    price_col = {site: 3 + i for i, site in enumerate(SITES)}
    link_col = {site: 11 + i for i, site in enumerate(SITES)}
    for index, row in enumerate(rows, start=2):
        _write_text(ws.cell(row=index, column=1), row.get("name"))
        _write_text(ws.cell(row=index, column=2), row.get("brand") or "")
        comparable = row.get("min_price") is not None
        differs = comparable and row["min_price"] != row["max_price"]
        for site in SITES:
            entry = row["prices"].get(site)
            if entry is None:
                continue
            cell = ws.cell(row=index, column=price_col[site], value=round(entry["price"], 2))
            cell.number_format = _MONEY
            if entry.get("stale"):
                cell.font = _STALE_FONT
            elif differs:
                value = _compared_value(row, entry)
                if value == row["min_price"]:
                    cell.font = _LOW_FONT
                elif value == row["max_price"]:
                    cell.font = _HIGH_FONT
            url = entry.get("url") or ""
            if url.startswith(("http://", "https://")) and not ILLEGAL_CHARACTERS_RE.search(url):
                link = ws.cell(row=index, column=link_col[site], value=SITE_TITLES[site])
                link.hyperlink = url
                link.font = _LINK_FONT
        _write_text(ws.cell(row=index, column=6), _cheapest_sites(row))
        if comparable:
            diff = ws.cell(row=index, column=7, value=round(row["max_price"] - row["min_price"], 2))
            diff.number_format = _MONEY
            pct = ws.cell(row=index, column=8, value=(row.get("spread_pct") or 0.0) / 100)
            pct.number_format = _PERCENT
        _write_text(ws.cell(row=index, column=9), _position(row, t))
        _write_text(ws.cell(row=index, column=10), _note(row, t))

    widths = [52, 22, 14, 14, 14, 24, 12, 10, 20, 60, 16, 16, 16]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"


def build_workbook(
    rows: list[dict[str, Any]],
    *,
    locale: str,
    generated_at: datetime,
    search: str | None = None,
    category: str | None = None,
    min_sites: int = 2,
    with_aloe: bool = False,
    diff_only: bool = False,
) -> bytes:
    """Собрать .xlsx из строк сравнения (формат `ComparisonRowOut`).

    `diff_only` — страница уже показывает только различия: второй лист тогда
    повторял бы первый, и его не пишем.
    """
    t = strings(locale)
    differing = [row for row in rows if has_price_difference(row)]

    wb = Workbook()
    ws_diff = wb.active
    ws_diff.title = t["sheet_diff"]
    _write_sheet(ws_diff, differing, t)
    if not diff_only:
        _write_sheet(wb.create_sheet(t["sheet_all"]), rows, t)

    info = wb.create_sheet(t["sheet_info"])
    local = generated_at + _BAKU_OFFSET
    lines: list[tuple[str, Any]] = [
        (t["info_title"], ""),
        (t["info_generated"], f"{local:%d.%m.%Y %H:%M} (Bakı)"),
        (t["info_search"], search or t["all"]),
        (t["info_category"], category or t["all"]),
        (t["info_min_sites"], min_sites),
    ]
    if with_aloe:
        lines.append((t["info_aloe"], t["yes"]))
    lines.append((t["info_diff_rows"], len(differing)))
    if not diff_only:
        lines.append((t["info_all_rows"], len(rows)))
    lines += [("", ""), (t["info_rule"], ""), (t["info_colors"], "")]
    for index, (label, value) in enumerate(lines, start=1):
        _write_text(info.cell(row=index, column=1), label)
        if isinstance(value, int):
            info.cell(row=index, column=2, value=value)
        else:
            _write_text(info.cell(row=index, column=2), value)
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 44
    info.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
