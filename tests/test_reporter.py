"""Тесты генерации HTML и Excel отчёта на синтетических данных."""

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from src import reporter
from src.analyzer import (
    AnalysisReport,
    CompetitorUndercut,
    NewProduct,
    PriceChange,
    PromoChange,
)


def _sample_report() -> AnalysisReport:
    return AnalysisReport(
        run_id=1,
        run_started_at=datetime(2026, 4, 28, 6, 0),
        prev_run_id=0,
        prev_run_at=datetime(2026, 4, 27, 6, 0),
        price_changes=[
            PriceChange(
                product_name="Test Pill 500mg",
                site="aloe",
                url="http://aloe.az/p/1",
                prev_price=10.0,
                curr_price=8.5,
                delta_pct=-15.0,
                is_on_sale=True,
                promo_label="-15%",
            ),
        ],
        undercuts=[
            CompetitorUndercut(
                canonical_name="Test Pill 500mg",
                client_price=10.0,
                client_url="http://pharmonline.az/p/1",
                competitor_site="aloe",
                competitor_price=8.5,
                competitor_url="http://aloe.az/p/1",
                diff_pct=15.0,
            ),
        ],
        new_products=[
            NewProduct(
                site="aptekonline",
                name="New SKU",
                url="http://aptekonline.az/p/2",
                price=12.0,
                category="vitamins",
            ),
        ],
        promo_changes=[
            PromoChange(
                site="aloe",
                title="Summer sale 30% off",
                landing_url="http://aloe.az/promo",
                is_new=True,
            ),
        ],
        summary_counts={
            "price_changes": 1,
            "undercuts": 1,
            "new_products": 1,
            "promo_changes": 1,
        },
    )


def test_render_html_contains_key_sections():
    html = reporter.render_html(_sample_report())
    assert "Test Pill 500mg" in html
    assert "Конкурент опустил цену" in html
    assert "изменения цен" in html.lower()
    assert "Новые товары" in html
    assert "−15.0%" in html or "-15.0%" in html


def test_render_excel_has_all_sheets():
    xlsx_bytes = reporter.render_excel(_sample_report())
    wb = load_workbook(BytesIO(xlsx_bytes))
    expected = {"Конкурент дешевле", "Изменения цен", "Новые SKU", "Промо"}
    actual = set(wb.sheetnames)
    assert expected.issubset(actual)


def test_email_subject_includes_alerts():
    subj = reporter.email_subject(_sample_report())
    assert "1 undercuts" in subj
    assert "Pharmacy Monitor" in subj
