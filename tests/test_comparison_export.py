"""Оформление Excel-выгрузки сравнения (src/comparison_export.py)."""

import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from src import comparison_export
from src.comparison_export import build_workbook, has_price_difference


def _row(name, prices, *, basis="raw", brand=None, needs_review=False):
    """Строка в формате `ComparisonRowOut`; min/max считаем как API — по свежим ценам."""
    key = "unit_price" if basis == "unit" else "price"
    fresh = [entry[key] for entry in prices.values() if not entry.get("stale")]
    comparable = len(fresh) >= 2
    low, high = (min(fresh), max(fresh)) if comparable else (None, None)
    return {
        "canonical_id": 1,
        "name": name,
        "brand": brand,
        "prices": prices,
        "min_price": low,
        "max_price": high,
        "spread_pct": round((high - low) / high * 100, 1) if comparable else None,
        "spread_basis": basis,
        "needs_review": needs_review,
    }


def _price(price, **extra):
    return {"price": price, "url": "https://site.example/p", **extra}


def _open(rows, **kwargs):
    kwargs.setdefault("locale", "ru")
    content = build_workbook(rows, generated_at=datetime(2026, 10, 6, 21, 30), **kwargs)
    return load_workbook(io.BytesIO(content))


def _body(ws):
    return [row for row in ws.iter_rows(min_row=2, values_only=True)]


def test_has_price_difference():
    assert has_price_difference(_row("a", {"pharmonline": _price(10.0), "aloe": _price(10.03)}))
    assert not has_price_difference(_row("b", {"pharmonline": _price(5.0), "aloe": _price(5.0)}))
    # Одна цена — сравнивать не с чем.
    assert not has_price_difference(_row("c", {"pharmonline": _price(5.0)}))


def test_filename_uses_baku_date():
    # 21:30 UTC — в Баку уже следующий день.
    assert comparison_export.filename("ru", datetime(2026, 10, 6, 21, 30)) == (
        "raznye-ceny-2026-10-07.xlsx"
    )
    assert comparison_export.filename("xx", datetime(2026, 10, 6, 9, 0)) == (
        "qiymet-ferqleri-2026-10-06.xlsx"
    )


@pytest.mark.parametrize(
    ("prices", "position", "cheapest"),
    [
        ({"pharmonline": _price(8.0), "aptekonline": _price(10.0)}, "Самая низкая", "Pharmonline"),
        ({"pharmonline": _price(10.0), "aloe": _price(8.0)}, "Самая высокая", "Aloe"),
        (
            {"pharmonline": _price(9.0), "aptekonline": _price(8.0), "aloe": _price(10.0)},
            "Посередине",
            "Aptekonline",
        ),
        (
            {"pharmonline": _price(8.0), "aptekonline": _price(8.0), "aloe": _price(10.0)},
            "Самая низкая",
            "Pharmonline, Aptekonline",
        ),
        ({"aptekonline": _price(8.0), "aloe": _price(10.0)}, None, "Aptekonline"),
    ],
)
def test_position_and_cheapest_columns(prices, position, cheapest):
    row = _body(_open([_row("Товар", prices)]).worksheets[0])[0]
    assert row[5] == cheapest
    assert row[8] == position


def test_prices_are_numbers_and_difference_is_a_percentage():
    ws = _open([_row("Товар", {"pharmonline": _price(10.0), "aptekonline": _price(7.5)})])[
        "Разные цены"
    ]
    assert ws["C2"].value == 10.0 and ws["C2"].number_format == "0.00"
    assert ws["G2"].value == 2.5
    assert ws["H2"].value == pytest.approx(0.25) and ws["H2"].number_format == "0.0%"
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:M2"


def test_unit_basis_and_stale_price_are_explained_in_the_note():
    rows = [
        _row(
            "Маска",
            {
                "pharmonline": _price(10.0, unit_price=0.2, pack_count=50),
                "aloe": _price(0.25, unit_price=0.25, pack_count=1),
                "aptekonline": _price(9.0, unit_price=9.0, stale=True, age_days=20),
            },
            basis="unit",
            needs_review=True,
        )
    ]
    row = _body(_open(rows).worksheets[0])[0]

    assert row[2:5] == (10.0, 9.0, 0.25)  # в колонках — цены упаковок, как на сайтах
    assert row[6] == pytest.approx(0.05)  # разница — за 1 шт.
    assert row[8] == "Самая низкая"
    note = row[9]
    assert "Pharmonline 0.20 (50 шт.)" in note and "Aloe 0.25 (1 шт.)" in note
    assert "Aptekonline: цене 20 дн." in note
    assert "проверьте, что товары одинаковые" in note


def test_scraped_text_cannot_become_a_formula_or_break_the_file():
    rows = [
        _row(
            '=HYPERLINK("http://evil.example","x")',
            {"pharmonline": _price(1.0), "aloe": _price(2.0)},
            brand="+SUM(1;1)\x07",
        )
    ]
    ws = _open(rows, search="@cmd").worksheets[0]

    assert ws["A2"].data_type == "s" and ws["A2"].value.startswith("=HYPERLINK")
    assert ws["B2"].data_type == "s" and ws["B2"].value == "+SUM(1;1)"
    info = _open(rows, search="@cmd")["О файле"]
    assert info["B3"].data_type == "s" and info["B3"].value == "@cmd"


def test_only_http_links_become_hyperlinks():
    rows = [
        _row(
            "Товар",
            {
                "pharmonline": {"price": 1.0, "url": "https://pharmonline.az/p/1"},
                "aloe": {"price": 2.0, "url": "javascript:alert(1)"},
            },
        )
    ]
    ws = _open(rows).worksheets[0]
    assert ws["K2"].hyperlink.target == "https://pharmonline.az/p/1"
    assert ws["M2"].value is None and ws["M2"].hyperlink is None


def test_sheets_and_info(tmp_path):
    rows = [
        _row("Разная", {"pharmonline": _price(10.0), "aloe": _price(8.0)}),
        _row("Одинаковая", {"pharmonline": _price(5.0), "aloe": _price(5.0)}),
    ]
    wb = _open(rows, search="крем", category="skin", min_sites=2, with_aloe=True)

    assert wb.sheetnames == ["Разные цены", "Всё сравнение", "О файле"]
    assert [r[0] for r in _body(wb["Разные цены"])] == ["Разная"]
    assert [r[0] for r in _body(wb["Всё сравнение"])] == ["Разная", "Одинаковая"]
    info = {row[0]: row[1] for row in wb["О файле"].iter_rows(values_only=True) if row[0]}
    assert info["Файл сформирован"] == "07.10.2026 01:30 (Bakı)"
    assert info["Поиск"] == "крем"
    assert info["Товаров с разной ценой"] == 1
    assert info["Всего товаров в выборке"] == 2

    assert _open(rows, diff_only=True).sheetnames == ["Разные цены", "О файле"]
    # Пустая выборка — валидный файл с заголовками, а не ошибка.
    assert _body(_open([]).worksheets[0]) == []
