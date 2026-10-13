"""Утреннее письмо о разнице цен — суточный дайджест с отчётом и Excel-вложением.

Клиент попросил присылать «различающиеся данные» на почту каждый день сам.
Письмо собирает `notifications.send_daily_digest`, шлёт его команда
`notify digest daily`. Что здесь проверяется:

- числа в письме можно пересчитать по приложенному файлу, а сам файл — тот, что
  отдаёт кнопка «Excel» на странице сравнения;
- письмо уходит и без событий за сутки; без отчёта и без событий — не уходит;
- сломанный отчёт не отменяет письмо с событиями, а когда слать больше нечего —
  команда кончается ошибкой, а не тихим пропуском;
- недельный дайджест отчёта не знает (его письмо закреплено байт в байт в
  tests/test_notifications.py).

Письма не уходят никуда: `notifier.send_email` подменён.
"""

from __future__ import annotations

import io
import os
import re
import uuid
from collections.abc import Iterator
from datetime import datetime, timedelta
from html import unescape
from unittest.mock import patch

import pytest
from click.testing import CliRunner
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, text
from sqlalchemy.orm import Session, sessionmaker
from structlog.testing import capture_logs

from src import api as api_module
from src import (
    cadence,
    comparison_export,
    comparison_rows,
    daily_letter_state,
    notifications,
    storage,
    tenants,
)
from src._time import utcnow
from src.main import cli
from src.storage import Base

NBSP = " "
# 05:00 UTC — время таймера: в Баку 09:00 того же дня.
MORNING = datetime(2026, 10, 13, 5, 0)
SITES = ("pharmonline", "aptekonline", "aloe")


@pytest.fixture
def setup(monkeypatch, db_session):
    engine = db_session.get_bind()
    Session = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(storage, "make_session", lambda *a, **kw: Session)
    monkeypatch.setattr(storage, "init_db", lambda *a, **kw: None)
    monkeypatch.setenv("PHARMACY_PUBLIC_URL", "https://example.test")
    # Политики стран и наличия — в shadow, как на проде: сравнение открыто.
    monkeypatch.delenv("COUNTRY_IDENTITY_POLICY", raising=False)
    monkeypatch.delenv("OFFER_AVAILABILITY_POLICY", raising=False)
    return db_session


def _reader(s, **fields) -> storage.TenantUser:
    """Сотрудник клиента с включённым суточным дайджестом."""
    tenant = tenants.get_or_create_default(s)
    defaults = {
        "email": "reader@example.test",
        "role": "viewer",
        "is_active": True,
        "daily_digest": True,
    }
    user = storage.TenantUser(tenant_id=tenant.id, created_at=utcnow(), **{**defaults, **fields})
    s.add(user)
    s.commit()
    return user


@pytest.fixture
def reader(setup) -> storage.TenantUser:
    return _reader(setup)


def _trusted_run(
    s, *, sites=SITES, finished_at: datetime | None = None, tenant_id: int = 1
) -> storage.Run:
    """Полный сбор, которому система верит: проверка каталога пройдена."""
    finished_at = finished_at or utcnow()
    run = storage.Run(
        tenant_id=tenant_id,
        status="ok",
        started_at=finished_at - timedelta(minutes=30),
        finished_at=finished_at,
        catalog_scope="full",
        full_catalog_sites=",".join(sites),
        catalog_verified=True,
        run_quality={
            "baseline_enforced": True,
            "full_catalog_verified": True,
            "financially_eligible": True,
            "sites": {site: {"status": "ok"} for site in sites},
        },
    )
    s.add(run)
    s.flush()
    return run


def _pair(s, run, name: str, prices: dict[str, float], *, stale=(), tenant_id: int = 1) -> None:
    """Пара товаров на сайтах с ценами; `stale` — сайты, где товар давно не видели."""
    match = storage.Match(tenant_id=tenant_id, canonical_name=name, confidence=1.0)
    s.add(match)
    s.flush()
    for site, price in prices.items():
        product = storage.Product(
            tenant_id=tenant_id,
            site=site,
            external_id=f"{site}-{name}-{tenant_id}",
            url=f"https://{site}.example/{name}",
            name=name,
            name_normalized=name.lower(),
            canonical_id=match.id,
            last_seen_at=utcnow() - timedelta(days=20 if site in stale else 0),
        )
        s.add(product)
        s.flush()
        s.add(
            storage.PriceSnapshot(
                run_id=run.id, product_id=product.id, price=price, captured_at=utcnow()
            )
        )
    s.commit()


def _mixed_catalog(s, run) -> None:
    """По товару на каждый исход, который письмо считает отдельно."""
    _pair(s, run, "Дороже всех", {"pharmonline": 10.0, "aloe": 8.0})
    _pair(s, run, "Посередине", {"pharmonline": 9.0, "aptekonline": 8.0, "aloe": 10.0})
    _pair(s, run, "Дешевле всех", {"pharmonline": 8.0, "aptekonline": 10.0})
    _pair(s, run, "Поровну с конкурентом", {"pharmonline": 8.0, "aptekonline": 8.0, "aloe": 9.0})
    _pair(s, run, "Только конкуренты", {"aptekonline": 8.0, "aloe": 10.0})
    _pair(
        s,
        run,
        "Цена клиента устарела",
        {"pharmonline": 7.0, "aptekonline": 8.0, "aloe": 9.0},
        stale=("pharmonline",),
    )
    _pair(s, run, "Одинаковая", {"pharmonline": 5.0, "aloe": 5.0})
    _pair(s, run, "Сравнить не с чем", {"pharmonline": 5.0, "aloe": 9.0}, stale=("aloe",))
    # На одном сайте — в сравнение не входит вовсе.
    _pair(s, run, "Один сайт", {"pharmonline": 3.0})


def _event(s, run, title: str = "Kreon 25000", severity: str = "critical") -> None:
    s.add(
        storage.AlertEvent(
            rule_type="undercut_threshold",
            dedup_key=f"letter|{title}",
            severity=severity,
            title=title,
            detail=f"detail of {title}",
            tenant_id=run.tenant_id,
            created_at=utcnow(),
            payload={"site": "aloe", "source_run_id": run.id},
        )
    )
    s.commit()


def _mail_server(accepts: bool = True):
    """Подмена `notifier.send_email`: почтовый сервер письмо принимает (или не настроен)."""
    return patch("src.notifier.send_email", return_value=accepts)


def _send_daily(s, **kwargs):
    """Отправить суточное письмо: (число получателей, вызовы `send_email`)."""
    with _mail_server() as mock_email:
        sent = notifications.send_daily_digest(s, tenant_id=1, **kwargs)
    return sent, mock_email.call_args_list


def _book(call):
    """Вложение письма, открытое как книга, и его имя."""
    (attachment,) = call.kwargs["attachments"]
    file_name, content, media_type = attachment
    assert media_type == comparison_export.XLSX_MEDIA_TYPE
    return load_workbook(io.BytesIO(content)), file_name


def _body_rows(ws) -> list[tuple]:
    return list(ws.iter_rows(min_row=2, values_only=True))


def _text(html: str) -> str:
    """Письмо так, как его читает человек: без разметки, пробелы схлопнуты."""
    return " ".join(unescape(re.sub(r"<[^>]+>", " ", html)).split())


def _number(letter_text: str, label: str) -> int:
    """Число, стоящее в письме сразу за подписью."""
    found = re.search(re.escape(label) + r" ([\d ]+)", letter_text.replace(NBSP, ""))
    assert found, f"в письме нет строки «{label}»: {letter_text}"
    return int(found.group(1).replace(" ", ""))


# ─── Числа письма и приложенный файл ─────────────────────────────────────────


def test_letter_numbers_can_be_recounted_from_the_attached_file(setup, reader):
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    sent, calls = _send_daily(s)

    assert sent == 1 and len(calls) == 1
    letter = _text(calls[0].kwargs["html_body"])
    book, _name = _book(calls[0])
    differing = _body_rows(book.worksheets[0])
    everything = _body_rows(book.worksheets[1])
    positions = [row[8] for row in differing]

    # «Различаются» в письме — это строки первого листа, и то же число в теме.
    assert _number(letter, "Цена на сайтах различается") == len(differing) == 6
    assert calls[0].kwargs["subject"].endswith(": различаются у 6 товаров")
    assert _number(letter, "Товаров в сравнении") == len(everything) == 8
    # Разбивка — колонка «Позиция Pharmonline» того же листа.
    assert _number(letter, "у конкурента дешевле, чем на Pharmonline") == (
        positions.count("Самая высокая") + positions.count("Посередине")
    )
    assert _number(letter, "на Pharmonline цена самая низкая") == positions.count("Самая низкая")
    assert _number(letter, "свежей цены Pharmonline нет — различаются цены конкурентов") == (
        positions.count(None)
    )
    assert _number(letter, "Цена одинаковая") == (
        [row[8] for row in everything].count("Цена одинаковая")
    )
    assert _number(letter, "Сравнить не с чем: свежая цена только на одном сайте") == 1
    # И сами значения — чтобы совпадение не было совпадением двух нулей.
    assert (positions.count("Самая низкая"), positions.count(None)) == (2, 2)
    assert positions.count("Самая высокая") + positions.count("Посередине") == 2
    assert {row[0] for row in differing} == {
        "Дороже всех",
        "Посередине",
        "Дешевле всех",
        "Поровну с конкурентом",
        "Только конкуренты",
        "Цена клиента устарела",
    }


def test_a_stale_price_is_not_a_difference(setup, reader):
    """Разной была бы только цена, которую сайт давно не показывает."""
    s = setup
    run = _trusted_run(s)
    _pair(
        s,
        run,
        "Разница только с устаревшей",
        {"pharmonline": 5.0, "aptekonline": 5.0, "aloe": 9.0},
        stale=("aloe",),
    )

    _sent, calls = _send_daily(s)

    letter = _text(calls[0].kwargs["html_body"])
    book, _name = _book(calls[0])
    assert _number(letter, "Цена на сайтах различается") == 0
    assert _number(letter, "Цена одинаковая") == 1
    assert _body_rows(book.worksheets[0]) == []
    assert calls[0].kwargs["subject"].endswith(": различий нет")


def test_letter_has_no_list_of_the_biggest_gaps(setup, reader):
    """Названий товаров в теле нет: наверху по разнице чаще стоят неверные пары."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    _sent, calls = _send_daily(s)

    assert "Дороже всех" not in calls[0].kwargs["html_body"]
    assert notifications._EVENT_ROW_MARK not in calls[0].kwargs["html_body"]


def test_attachment_is_the_file_of_the_comparison_page(setup, reader, monkeypatch):
    """Тот же файл, что отдаёт кнопка «Excel» на странице с фильтрами по умолчанию."""
    if not api_module._JWT_AVAILABLE:
        pytest.skip("python-jose not installed")
    monkeypatch.setenv("JWT_SECRET", "test-secret-very-long-not-for-prod-only")
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    admin = _reader(s, email="admin@example.test", role="admin", daily_digest=False)

    _sent, calls = _send_daily(s)
    attached, file_name = _book(calls[0])

    client = TestClient(api_module.app)
    client.get(f"/auth/verify?token={tenants.issue_magic_token(s, admin.email)}")
    response = client.get("/api/v1/dash/comparison/export.xlsx?locale=ru")
    assert response.status_code == 200, response.text
    downloaded = load_workbook(io.BytesIO(response.content))

    assert f'filename="{file_name}"' in response.headers["content-disposition"]
    assert (
        attached.sheetnames == downloaded.sheetnames == ["Разные цены", "Всё сравнение", "О файле"]
    )
    for ours, theirs in zip(attached.worksheets[:2], downloaded.worksheets[:2], strict=True):
        assert _body_rows(ours) == _body_rows(theirs)
        assert [cell.value for cell in ours[1]] == [cell.value for cell in theirs[1]]
    # Лист «О файле»: те же фильтры и счётчики; минута выгрузки у каждого своя.
    about = [
        {row[0]: row[1] for row in book["О файле"].iter_rows(values_only=True) if row[0]}
        for book in (attached, downloaded)
    ]
    for facts in about:
        facts.pop("Файл сформирован")
    assert about[0] == about[1]
    assert about[0]["Минимум сайтов"] == 2 and about[0]["Поиск"] == "все"


def test_attachment_name_and_date_are_bakus(setup, reader, monkeypatch):
    """21:30 UTC — в Баку уже следующий день: дата в теме, в письме и в имени файла одна."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    monkeypatch.setattr(notifications, "utcnow", lambda: datetime(2026, 10, 13, 21, 30))

    _sent, calls = _send_daily(s)

    book, file_name = _book(calls[0])
    assert file_name == "raznye-ceny-2026-10-14.xlsx"
    assert calls[0].kwargs["subject"].startswith("Pharmacy Monitor — цены на 14.10.2026: ")
    letter = _text(calls[0].kwargs["html_body"])
    assert "Разница цен на 14.10.2026" in letter
    assert f"в приложенном файле {file_name}" in letter
    about = {row[0]: row[1] for row in book["О файле"].iter_rows(values_only=True) if row[0]}
    assert about["Файл сформирован"] == "14.10.2026 01:30 (Bakı)"


def test_mail_server_receives_the_letter_with_the_workbook_attached(setup, reader, monkeypatch):
    """Без подмены `send_email`: до почтового сервера доходит письмо с файлом, который открывается."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    delivered = []

    class Server:
        def __init__(self, *args, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def starttls(self):
            pass

        def login(self, *args):
            pass

        def send_message(self, message):
            delivered.append(message)

    monkeypatch.setattr("smtplib.SMTP", Server)
    monkeypatch.setenv("SMTP_HOST", "smtp.example.test")
    monkeypatch.setenv("SMTP_USER", "sender")
    monkeypatch.setenv("SMTP_PASSWORD", "not-a-secret")
    monkeypatch.setenv("SMTP_FROM", "Pharmacy Monitor")

    assert notifications.send_daily_digest(s, tenant_id=1, locale="az") == 1

    (message,) = delivered
    assert message["To"] == "reader@example.test"
    assert message["Subject"].startswith("Pharmacy Monitor — цены на ")
    assert "Разница цен на" in message.get_body(preferencelist=("html",)).get_content()
    (attached,) = message.iter_attachments()
    assert attached.get_content_type() == comparison_export.XLSX_MEDIA_TYPE
    assert re.fullmatch(r"qiymet-ferqleri-\d{4}-\d{2}-\d{2}\.xlsx", attached.get_filename())
    book = load_workbook(io.BytesIO(attached.get_content()))
    assert len(_body_rows(book.worksheets[0])) == 6


def test_locale_changes_the_attachment_and_not_the_letter(setup, reader):
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    _sent, calls = _send_daily(s, locale="az")

    book, file_name = _book(calls[0])
    assert file_name.startswith("qiymet-ferqleri-")
    assert book.sheetnames == ["Fərqli qiymətlər", "Bütün müqayisə", "Məlumat"]
    assert _body_rows(book.worksheets[0])[0][8] in {"Ən baha", "Ən ucuz", "Ortada", None}
    # Тело — русское, как у всех писем системы.
    letter = _text(calls[0].kwargs["html_body"])
    assert "Разница цен на" in letter and "Fərqli" not in letter
    assert calls[0].kwargs["subject"].startswith("Pharmacy Monitor — цены на ")


def test_another_tenants_pairs_are_not_counted(setup, reader):
    s = setup
    _pair(s, _trusted_run(s), "Свой товар", {"pharmonline": 10.0, "aloe": 8.0})
    s.add(storage.Tenant(id=2, slug="other", name="Other", client_site="pharmonline"))
    _pair(
        s,
        _trusted_run(s, tenant_id=2),
        "Чужой товар",
        {"pharmonline": 1.0, "aloe": 2.0},
        tenant_id=2,
    )

    _sent, calls = _send_daily(s)

    book, _name = _book(calls[0])
    assert [row[0] for row in _body_rows(book.worksheets[1])] == ["Свой товар"]
    assert _number(_text(calls[0].kwargs["html_body"]), "Товаров в сравнении") == 1


# ─── Когда письмо уходит ─────────────────────────────────────────────────────


def test_letter_goes_out_on_a_morning_without_events(setup, reader):
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    sent, calls = _send_daily(s)

    assert sent == 1
    assert calls[0].kwargs["to"] == ["reader@example.test"]
    letter = _text(calls[0].kwargs["html_body"])
    assert notifications._NO_EVENTS_LINE in letter
    # Сводки событий нет вовсе — ни заголовка, ни «0 событий».
    assert "События за сутки" not in letter and "Дайджест за сутки" not in letter
    assert "0 событий" not in letter


def test_letter_with_events_puts_the_price_report_above_them(setup, reader):
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run, "Kreon 25000")

    _sent, calls = _send_daily(s)

    html = calls[0].kwargs["html_body"]
    letter = _text(html)
    assert letter.index("Разница цен на") < letter.index("События за сутки")
    assert letter.index("События за сутки") < letter.index("Kreon 25000")
    assert notifications._NO_EVENTS_LINE not in letter
    assert html.count(notifications._EVENT_ROW_MARK) == 1
    assert 'href="https://example.test/alerts?hours=24"' in html
    assert 'href="https://example.test/comparison?diff=1"' in html
    assert len(calls[0].kwargs["attachments"]) == 1


def test_nothing_is_sent_with_neither_a_report_nor_events(setup, reader):
    """Сборов ещё не было, событий нет — как и раньше, письма нет."""
    with capture_logs() as logs:
        sent, calls = _send_daily(setup)

    assert sent == 0 and calls == []
    assert {"event": "digest_price_report_skipped", "reason": "no_successful_run"}.items() <= (
        logs[0].items()
    )
    assert [entry["event"] for entry in logs][-1] == "digest_no_events"


def _count_report_builds(monkeypatch) -> list[int]:
    """Сколько раз звали сборку строк сравнения. Считаем, а не роняем: исключение
    из сборки отчёта дайджест ловит сам, и письмо ушло бы как ни в чём не бывало."""
    built: list[int] = []
    real = comparison_rows.comparison_rows

    def counted(*args, **kwargs):
        built.append(1)
        return real(*args, **kwargs)

    monkeypatch.setattr(comparison_rows, "comparison_rows", counted)
    return built


def test_no_recipients_means_no_report_is_built(setup, monkeypatch):
    """Отчёт — тысячи строк и книга Excel: без получателей его не собирают."""
    s = setup
    _reader(s, daily_digest=False)
    _mixed_catalog(s, _trusted_run(s))
    built = _count_report_builds(monkeypatch)

    sent, calls = _send_daily(s)

    assert sent == 0 and calls == []
    assert built == []


def test_without_a_report_the_daily_letter_is_the_old_digest(setup, reader):
    """Сравнивать нечего (пар нет) — письмо о событиях прежнее: тема, тело, без вложения."""
    s = setup
    run = _trusted_run(s)
    _event(s, run, "Kreon 25000")

    with capture_logs() as logs:
        sent, calls = _send_daily(s)

    assert sent == 1
    assert set(calls[0].kwargs) == {"subject", "html_body", "to"}
    assert calls[0].kwargs["subject"] == (
        "Pharmacy Monitor — дайджест за сутки: 1 событие, из них 1 критичное"
    )
    letter = _text(calls[0].kwargs["html_body"])
    assert "Дайджест за сутки" in letter and "Разница цен" not in letter
    skipped = [entry for entry in logs if entry["event"] == "digest_price_report_skipped"]
    assert [entry["reason"] for entry in skipped] == ["no_comparison_rows"]


def test_paused_comparison_leaves_the_events_letter_without_the_report(setup, reader, monkeypatch):
    """Политика в enforce без подтверждённого каталога: страница отвечает 503, письмо — без отчёта."""
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run)
    monkeypatch.setenv("COUNTRY_IDENTITY_POLICY", "enforce")

    with capture_logs() as logs:
        sent, calls = _send_daily(s)

    assert sent == 1
    assert "attachments" not in calls[0].kwargs
    assert "Разница цен" not in calls[0].kwargs["html_body"]
    skipped = [entry for entry in logs if entry["event"] == "digest_price_report_skipped"]
    assert [entry["reason"] for entry in skipped] == ["full_catalog_trust_not_ready"]


# ─── Сломанный отчёт ─────────────────────────────────────────────────────────


def _broken(*args, **kwargs):
    raise KeyError("pack_size")


@pytest.mark.parametrize(
    ("module", "name"),
    [(comparison_rows, "comparison_rows"), (comparison_export, "build_workbook")],
)
def test_events_letter_survives_a_broken_report(setup, reader, monkeypatch, module, name):
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run, "Kreon 25000")
    monkeypatch.setattr(module, name, _broken)

    with capture_logs() as logs:
        sent, calls = _send_daily(s)

    assert sent == 1
    # Письмо о событиях — прежнее: без отчёта, без вложения, с прежней темой.
    assert set(calls[0].kwargs) == {"subject", "html_body", "to"}
    assert calls[0].kwargs["subject"].startswith("Pharmacy Monitor — дайджест за сутки: 1 событие")
    assert "Kreon 25000" in calls[0].kwargs["html_body"]
    assert "Разница цен" not in calls[0].kwargs["html_body"]
    # Причина — в журнале, без адреса.
    failed = [entry for entry in logs if entry["event"] == "digest_price_report_failed"]
    assert len(failed) == 1
    assert failed[0]["error_type"] == "KeyError" and failed[0]["log_level"] == "error"
    assert "@" not in repr(logs)


def test_broken_report_with_nothing_else_to_send_is_an_error(setup, reader, monkeypatch):
    """Тихий пропуск скрыл бы сломанное утреннее письмо."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    monkeypatch.setattr(comparison_rows, "comparison_rows", _broken)

    with patch("src.notifier.send_email") as mock_email:
        with pytest.raises(notifications.DailyLetterNotSent) as raised:
            notifications.send_daily_digest(s, tenant_id=1)

    assert not mock_email.called
    assert raised.value.error_type == "KeyError"
    assert isinstance(raised.value.__cause__, KeyError)
    # Текст печатают команда и ответ API: в нём класс ошибки, а не её текст.
    assert "KeyError" in str(raised.value) and "pack_size" not in str(raised.value)


@pytest.mark.parametrize("extra", [(), ("--dry-run",)])
def test_command_exits_non_zero_when_the_letter_could_not_be_built(
    setup, reader, monkeypatch, extra
):
    """Юнит таймера должен показать сбой: код возврата — не ноль."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    monkeypatch.setattr(comparison_rows, "comparison_rows", _broken)

    with _mail_server() as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily", *extra])

    assert result.exit_code == 1, result.output
    assert "Error: отчёт о разнице цен не собрался (KeyError)" in result.output
    assert "письмо не отправлено" in result.output
    assert not mock_email.called
    assert "@" not in result.output


def test_command_fails_after_sending_the_events_letter_without_the_report(
    setup, reader, monkeypatch
):
    """Письмо о событиях ушло, но без отчёта — утро не состоялось: код не ноль.
    Сначала отправка, потом ошибка."""
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run)
    monkeypatch.setattr(comparison_rows, "comparison_rows", _broken)

    with _mail_server() as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily"])

    assert mock_email.call_count == 1
    assert "attachments" not in mock_email.call_args.kwargs
    assert result.exit_code == 1, result.output
    assert "daily digest sent to 1 recipients" in result.output
    assert "OK:" not in result.output
    assert "Error: отчёт о разнице цен не собрался (KeyError) — письмо без чисел и без файла" in (
        result.output
    )
    # Исход записан до ошибки: сторож в health видит то же, что код возврата.
    assert daily_letter_state.last_mailing(1).report == "failed:KeyError"


def test_dry_run_with_a_broken_report_fails_even_when_there_are_events(setup, reader, monkeypatch):
    """Просмотр перед включением должен показать сломанный отчёт, а не «всё соберётся»."""
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run)
    monkeypatch.setattr(comparison_export, "build_workbook", _broken)

    with _mail_server() as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily", "--dry-run"])

    assert result.exit_code == 1, result.output
    assert "[dry-run] daily digest: 1 recipients, nothing sent" in result.output
    assert "отчёт о разнице цен не собрался (KeyError)" in result.output
    assert not mock_email.called
    assert daily_letter_state.last_mailing(1) is None


def _refusing(*args, **kwargs):
    raise RuntimeError("{'reader@example.test': (550, b'5.1.1 mailbox unavailable')}")


@pytest.mark.parametrize(
    "server",
    [
        pytest.param({"side_effect": _refusing}, id="server-refused"),
        pytest.param({"return_value": False}, id="smtp-not-configured"),
    ],
)
def test_command_fails_when_no_recipient_got_the_letter(setup, reader, server):
    """Письмо собрано, но почтовый сервер не принял его ни у кого — это не «OK»."""
    s = setup
    _reader(s, email="second@example.test")
    _mixed_catalog(s, _trusted_run(s))

    with patch("src.notifier.send_email", **server) as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily"])

    assert mock_email.call_count == 2  # отказ одному не отменяет попытку другому
    assert result.exit_code == 1, result.output
    assert "почтовый сервер не принял письмо ни у одного из 2 получателей" in result.output
    assert "@" not in result.output


def test_command_succeeds_when_at_least_one_recipient_got_the_letter(setup, reader):
    s = setup
    _reader(s, email="second@example.test")
    _mixed_catalog(s, _trusted_run(s))
    answers = iter([RuntimeError("refused"), True])

    def one_of_two(*args, **kwargs):
        answer = next(answers)
        if isinstance(answer, Exception):
            raise answer
        return answer

    with patch("src.notifier.send_email", side_effect=one_of_two):
        result = CliRunner().invoke(cli, ["notify", "digest", "daily"])

    assert result.exit_code == 0, result.output
    assert "OK: daily digest sent to 1 recipients" in result.output
    state = daily_letter_state.last_mailing(1)
    assert (state.recipients, state.accepted, state.report) == (2, 1, "ok")


def test_dry_run_is_not_an_undelivered_letter(setup, reader):
    """При просмотре сервер ничего не принимал по определению — это не отказ."""
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    with _mail_server() as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily", "--dry-run"])

    assert result.exit_code == 0, result.output
    assert not mock_email.called


def test_quiet_morning_without_report_or_events_is_not_an_error(setup, reader):
    """Сборов ещё не было и событий нет: слать нечего, и это не сбой."""
    with _mail_server() as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "daily"])

    assert result.exit_code == 0, result.output
    assert "OK: daily digest sent to 0 recipients" in result.output
    assert not mock_email.called


def test_weekly_command_exit_code_does_not_follow_the_new_rules(setup):
    """Недельный дайджест не менялся: отказ почтового сервера у него — по-прежнему код 0."""
    s = setup
    _reader(s, daily_digest=False, weekly_digest=True)
    run = _trusted_run(s)
    _event(s, run)

    with patch("src.notifier.send_email", side_effect=_refusing) as mock_email:
        result = CliRunner().invoke(cli, ["notify", "digest", "weekly"])

    assert mock_email.call_count == 1
    assert result.exit_code == 0, result.output
    assert "OK: weekly digest sent to 0 recipients" in result.output
    assert daily_letter_state.last_mailing(1) is None


def _database_error(session: Session, *, tenant_id: int):
    """Сборка строк, у которой упал запрос к базе."""
    session.execute(text("select * from no_such_table_for_the_letter"))


def test_failed_report_query_is_rolled_back_to_a_savepoint(setup, reader, monkeypatch):
    """Отчёт и события читает одна сессия, и упавший запрос отчёта откатывается к
    точке сохранения. Зачем она нужна, SQLite не покажет (упавший запрос транзакцию
    тут не прерывает) — это тест на PostgreSQL ниже; здесь видно, что она есть."""
    s = setup
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run, "Kreon 25000")
    monkeypatch.setattr(comparison_rows, "comparison_rows", _database_error)
    statements: list[str] = []
    engine = s.get_bind()

    def record(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", record)
    try:
        sent, calls = _send_daily(s)
    finally:
        event.remove(engine, "before_cursor_execute", record)

    assert sent == 1 and "Kreon 25000" in calls[0].kwargs["html_body"]
    failed_at = next(i for i, sql in enumerate(statements) if "no_such_table_for_the_letter" in sql)
    assert any(sql.startswith("SAVEPOINT ") for sql in statements[:failed_at])
    assert statements[failed_at + 1].startswith("ROLLBACK TO SAVEPOINT ")
    # События читаются уже после отката.
    assert any("alert_events" in sql for sql in statements[failed_at + 2 :])


# ─── То же на PostgreSQL: упавший запрос прерывает транзакцию ────────────────


def _postgres_url() -> str:
    url = os.environ.get("DATABASE_URL", "")
    if url.startswith("postgresql"):
        return url
    if os.environ.get("CI"):
        pytest.fail(
            "в CI этот тест обязан исполняться, а DATABASE_URL — не PostgreSQL: на SQLite "
            "упавший запрос транзакцию не прерывает, и точку сохранения она не проверит."
        )
    pytest.skip("PostgreSQL DATABASE_URL is required")


@pytest.fixture
def postgres_session(monkeypatch) -> Iterator[Session]:
    """Сессия как в `storage.make_session` на своей схеме PostgreSQL."""
    url = _postgres_url()
    name = f"daily_letter_{uuid.uuid4().hex[:12]}"
    admin = create_engine(url)
    with admin.begin() as connection:
        connection.execute(text(f'CREATE SCHEMA "{name}"'))
    engine = create_engine(url, connect_args={"options": f"-csearch_path={name}"})
    monkeypatch.setenv("PHARMACY_PUBLIC_URL", "https://example.test")
    monkeypatch.delenv("COUNTRY_IDENTITY_POLICY", raising=False)
    monkeypatch.delenv("OFFER_AVAILABILITY_POLICY", raising=False)
    try:
        Base.metadata.create_all(engine)
        with sessionmaker(engine, expire_on_commit=False, autoflush=False)() as session:
            yield session
    finally:
        engine.dispose()
        with admin.begin() as connection:
            connection.execute(text(f'DROP SCHEMA "{name}" CASCADE'))
        admin.dispose()


def test_failed_report_query_does_not_abort_the_digest_on_postgres(postgres_session, monkeypatch):
    """На PostgreSQL после упавшего запроса транзакция прервана: без точки
    сохранения чтение событий падало бы следом, и письмо не уходило бы вовсе."""
    s = postgres_session
    _reader(s)
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run, "Kreon 25000")
    monkeypatch.setattr(comparison_rows, "comparison_rows", _database_error)

    sent, calls = _send_daily(s)

    assert sent == 1 and "Kreon 25000" in calls[0].kwargs["html_body"]
    assert "attachments" not in calls[0].kwargs


def test_letter_is_built_on_postgres(postgres_session):
    """Сборка отчёта целиком на настоящей базе: числа, вложение, возраст данных."""
    s = postgres_session
    _reader(s)
    _mixed_catalog(s, _trusted_run(s))

    sent, calls = _send_daily(s, locale="az")

    assert sent == 1
    book, file_name = _book(calls[0])
    assert file_name.startswith("qiymet-ferqleri-")
    assert len(_body_rows(book.worksheets[0])) == 6
    assert _number(_text(calls[0].kwargs["html_body"]), "Цена на сайтах различается") == 6


# ─── От какого числа данные сайтов ───────────────────────────────────────────


def _hours(site: str) -> timedelta:
    return timedelta(hours=cadence.site_max_age_hours(site))


def test_letter_says_when_each_site_was_collected_and_warns_about_an_old_one(
    setup, reader, monkeypatch
):
    s = setup
    monkeypatch.setattr(notifications, "utcnow", lambda: MORNING)
    fresh = _trusted_run(
        s, sites=("pharmonline", "aptekonline"), finished_at=MORNING - timedelta(hours=4)
    )
    aloe_at = MORNING - _hours("aloe") - timedelta(days=3)
    _trusted_run(s, sites=("aloe",), finished_at=aloe_at)
    # Более поздний сбор aloe проверку не прошёл: данные по-прежнему от прошлого.
    s.add(
        storage.Run(
            tenant_id=1,
            status="degraded",
            started_at=MORNING - timedelta(hours=5),
            finished_at=MORNING - timedelta(hours=4),
            catalog_scope="full",
            full_catalog_sites="aloe",
            catalog_verified=False,
            run_quality={"financially_eligible": False, "sites": {"aloe": {"status": "failed"}}},
        )
    )
    _pair(s, fresh, "Товар", {"pharmonline": 10.0, "aptekonline": 8.0})

    with capture_logs() as logs:
        _sent, calls = _send_daily(s)

    letter = _text(calls[0].kwargs["html_body"])
    aloe_day = comparison_export.baku_time(aloe_at)
    # Время бакинское: 01:00 UTC — это 05:00.
    assert "Pharmonline 13.10.2026 05:00" in letter
    assert "Aptekonline 13.10.2026 05:00" in letter
    assert f"Aloe {aloe_day:%d.%m.%Y %H:%M}" in letter
    assert (
        f"Данные Aloe — от {aloe_day:%d.%m.%Y}: с тех пор полный сбор этого сайта не прошёл"
    ) in letter
    assert "Данные Pharmonline — от" not in letter and "Данные Aptekonline — от" not in letter
    report = next(entry for entry in logs if entry["event"] == "digest_price_report")
    assert report["overdue_sites"] == ["aloe"]


def test_site_data_is_old_only_past_what_its_rhythm_allows(setup, monkeypatch):
    """Порог — ритм сбора сайта с запасом (`cadence.site_max_age_hours`), не своё число."""
    s = setup
    tenants.get_or_create_default(s)
    monkeypatch.setattr(notifications, "utcnow", lambda: MORNING)
    _trusted_run(s, sites=("pharmonline",), finished_at=MORNING - _hours("pharmonline"))
    _trusted_run(
        s,
        sites=("aptekonline",),
        finished_at=MORNING - _hours("aptekonline") - timedelta(minutes=1),
    )
    s.commit()

    ages = {age.site: age for age in notifications._site_data_ages(s, 1, MORNING)}

    assert ages["pharmonline"].overdue is False
    assert ages["aptekonline"].overdue is True
    # У aloe проверенного полного сбора не было вовсе.
    assert ages["aloe"].collected_at is None and ages["aloe"].overdue is True


def test_old_data_warning_follows_the_rhythm_map(setup, reader, monkeypatch):
    """Ритм сбора меняют в одном месте — `cadence.SITE_SCRAPE_CADENCE_HOURS`. Те же
    данные двухдневной давности при недельном ритме в порядке, при суточном — старые."""
    s = setup
    monkeypatch.setattr(notifications, "utcnow", lambda: MORNING)
    run = _trusted_run(s, finished_at=MORNING - timedelta(days=2))
    _pair(s, run, "Товар", {"pharmonline": 10.0, "aptekonline": 8.0})

    def warned_sites() -> list[str]:
        _sent, calls = _send_daily(s)
        return re.findall(r"Данные (\w+) — от 11\.10\.2026", _text(calls[0].kwargs["html_body"]))

    for site in SITES:
        monkeypatch.setitem(cadence.SITE_SCRAPE_CADENCE_HOURS, site, 168)
    assert warned_sites() == []

    monkeypatch.setitem(cadence.SITE_SCRAPE_CADENCE_HOURS, "aloe", 24)
    assert cadence.site_max_age_hours("aloe") == 24 + cadence.CADENCE_GRACE_HOURS
    assert warned_sites() == ["Aloe"]

    for site in SITES:
        monkeypatch.setitem(cadence.SITE_SCRAPE_CADENCE_HOURS, site, 24)
    assert warned_sites() == ["Pharmonline", "Aptekonline", "Aloe"]


def test_one_missed_night_of_a_daily_rhythm_is_not_yet_old_data(setup, monkeypatch):
    """При суточном ритме порог — 30 часов (сутки и запас). Письмо уходит в 05:00
    UTC, ночной сбор кончается около 01:00: после одной пропущенной ночи данным
    28 часов — строки-предупреждения ещё нет, дата в письме вчерашняя. После двух
    ночей — есть."""
    s = setup
    tenants.get_or_create_default(s)
    monkeypatch.setitem(cadence.SITE_SCRAPE_CADENCE_HOURS, "aloe", 24)
    last_good_night = datetime(2026, 10, 12, 1, 0)
    _trusted_run(s, sites=("aloe",), finished_at=last_good_night)
    s.commit()

    def aloe(letter_at: datetime):
        return {age.site: age for age in notifications._site_data_ages(s, 1, letter_at)}["aloe"]

    assert aloe(datetime(2026, 10, 12, 5, 0)).overdue is False  # сбор был этой ночью
    assert aloe(datetime(2026, 10, 13, 5, 0)).overdue is False  # одну ночь пропустили: 28 ч
    assert aloe(datetime(2026, 10, 14, 5, 0)).overdue is True  # две ночи: 52 ч


def test_site_without_a_trusted_full_run_is_named_in_the_letter(setup, reader):
    s = setup
    run = _trusted_run(s, sites=("pharmonline", "aptekonline"))
    _pair(s, run, "Товар", {"pharmonline": 10.0, "aptekonline": 8.0})

    _sent, calls = _send_daily(s)

    letter = _text(calls[0].kwargs["html_body"])
    assert "Aloe: полного сбора, прошедшего проверку, ещё не было." in letter


def test_partial_run_does_not_count_as_collecting_a_site(setup, monkeypatch):
    """Тик освежает товары, но сайт целиком не собирает: дата остаётся от полного сбора."""
    s = setup
    tenants.get_or_create_default(s)
    full_at = MORNING - timedelta(days=2)
    _trusted_run(s, sites=("aloe",), finished_at=full_at)
    s.add(
        storage.Run(
            tenant_id=1,
            status="ok",
            started_at=MORNING - timedelta(hours=2),
            finished_at=MORNING - timedelta(hours=1),
            catalog_scope="partial",
            run_quality={"sites": {"aloe": {"status": "ok"}}},
        )
    )
    s.commit()

    ages = {age.site: age for age in notifications._site_data_ages(s, 1, MORNING)}

    assert ages["aloe"].collected_at == full_at


# ─── Просмотр без отправки, команда, кнопка в дашборде ───────────────────────


def test_dry_run_sends_nothing_and_logs_the_numbers_and_the_file_size(setup, reader):
    s = setup
    _mixed_catalog(s, _trusted_run(s))

    with capture_logs() as logs:
        sent, calls = _send_daily(s, dry_run=True)

    assert sent == 1 and calls == []
    by_event = {entry["event"]: entry for entry in logs}
    report = by_event["digest_price_report"]
    assert {
        key: report[key]
        for key in ("compared", "differing", "competitor_cheaper", "client_cheapest")
    } == {"compared": 8, "differing": 6, "competitor_cheaper": 2, "client_cheapest": 2}
    assert report["no_client_price"] == 2 and report["locale"] == "ru"
    assert report["attachment_kb"] > 1
    dry_run = by_event["digest_dry_run"]
    assert dry_run["user_id"] == reader.id
    assert dry_run["attachment_kb"] == report["attachment_kb"]
    assert dry_run["subject"].endswith(": различаются у 6 товаров")
    assert by_event["digest_dry_run_done"]["recipients"] == 1
    assert "@" not in repr(logs)


def test_command_passes_the_locale_and_keeps_only_and_dry_run(setup, reader):
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    runner = CliRunner()

    with _mail_server() as mock_email:
        result = runner.invoke(cli, ["notify", "digest", "daily", "--locale", "az"])
        assert result.exit_code == 0, result.output
        assert "OK: daily digest sent to 1 recipients" in result.output
        assert _book(mock_email.call_args)[1].startswith("qiymet-ferqleri-")

        # Без флага — русский файл.
        runner.invoke(cli, ["notify", "digest", "daily"])
        assert _book(mock_email.call_args)[1].startswith("raznye-ceny-")
        assert mock_email.call_count == 2

        result = runner.invoke(cli, ["notify", "digest", "daily", "--locale", "az", "--dry-run"])
        assert result.exit_code == 0, result.output
        assert "1 recipients, nothing sent" in result.output

        result = runner.invoke(cli, ["notify", "digest", "daily", "--only", "nobody@example.test"])
        assert result.exit_code == 0 and "sent to 0 recipients" in result.output
        assert mock_email.call_count == 2

    result = runner.invoke(cli, ["notify", "digest", "daily", "--locale", "de"])
    assert result.exit_code == 2 and "'de' is not one of" in result.output


def test_weekly_digest_never_builds_the_price_report(setup, monkeypatch):
    """У недельного письма ни отчёта, ни вложения — даже когда сравнение есть."""
    s = setup
    _reader(s, daily_digest=False, weekly_digest=True)
    run = _trusted_run(s)
    _mixed_catalog(s, run)
    _event(s, run, "Kreon 25000")
    built = _count_report_builds(monkeypatch)

    with patch("src.notifier.send_email") as mock_email:
        assert notifications.send_weekly_digest(s, tenant_id=1) == 1
        result = CliRunner().invoke(cli, ["notify", "digest", "weekly", "--locale", "az"])

    assert result.exit_code == 0, result.output
    assert built == []
    assert mock_email.call_count == 2
    for call in mock_email.call_args_list:
        assert set(call.kwargs) == {"subject", "html_body", "to"}
        assert call.kwargs["subject"].startswith("Pharmacy Monitor — дайджест за неделю: ")
        assert "Разница цен" not in call.kwargs["html_body"]


def test_dashboard_button_sends_the_same_letter(setup, reader, monkeypatch):
    """«Отправить дайджест» в Quick Actions зовёт ту же функцию: письмо то же, файл русский."""
    if not api_module._JWT_AVAILABLE:
        pytest.skip("python-jose not installed")
    monkeypatch.setenv("JWT_SECRET", "test-secret-very-long-not-for-prod-only")
    s = setup
    _mixed_catalog(s, _trusted_run(s))
    admin = _reader(s, email="admin@example.test", role="admin", daily_digest=False)
    client = TestClient(api_module.app)
    client.get(f"/auth/verify?token={tenants.issue_magic_token(s, admin.email)}")

    with patch("src.notifier.send_email") as mock_email:
        response = client.post("/api/v1/dash/digest/send-test?kind=daily")

    assert response.status_code == 200, response.text
    assert response.json() == {"ok": True, "recipients_sent": 1}
    assert mock_email.call_args.kwargs["to"] == ["reader@example.test"]
    assert _book(mock_email.call_args)[1].startswith("raznye-ceny-")
